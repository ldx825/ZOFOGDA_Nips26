#!/usr/bin/env python3
"""Refresh the rebuttal Excel result workbook and query-budget plots."""
from __future__ import annotations
import argparse
import json
import math
import time
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
REB = ROOT / "rebuttal"
REGISTRY = REB / "tables" / "experiment_registry.csv"
WORKBOOK = REB / "tables" / "rebuttal_experiment_results.xlsx"
PLOTS = REB / "plots"


def collect():
    registry = pd.read_csv(REGISTRY)
    result_rows = []
    curves = []
    status_map = {}
    for _, reg in registry.iterrows():
        out = ROOT / str(reg["output_dir"])
        summary_path = out / "summary.json"
        status = "pending"
        summary = {}
        if summary_path.exists():
            try:
                summary = json.loads(summary_path.read_text(encoding="utf-8"))
                status = summary.get("status", "unknown")
            except Exception as exc:
                status = "broken"
                summary = {"error": repr(exc)}
        status_map[reg["run_id"]] = status
        row = reg.to_dict()
        row["status"] = status
        for k, v in summary.items():
            if k not in row or k in {
                "final_test_metric", "best_test_metric", "final_train_objective", "global_steps",
                "cum_train_forward_queries", "cum_x_function_queries", "cum_train_backward_calls",
                "cum_train_sample_forwards", "train_wall_time_sec", "eval_wall_time_sec",
                "n_train", "n_test", "feature_dim", "parameter_dim", "completed_at", "error"
            }:
                row[k] = v
        result_rows.append(row)
        metrics = out / "metrics.csv"
        if metrics.exists():
            try:
                df = pd.read_csv(metrics)
                if not df.empty:
                    df["experiment_groups"] = reg["experiment_groups"]
                    df["notes"] = reg.get("notes", "")
                    df["batch_size"] = reg["batch_size"]
                    df["hidden"] = reg["hidden"]
                    df["c1"] = reg["c1"]
                    df["c2"] = reg["c2"]
                    df["direction_distribution"] = reg["direction_distribution"]
                    curves.append(df)
            except Exception:
                pass
    registry["status"] = registry["run_id"].map(status_map)
    raw = pd.DataFrame(result_rows)
    curve_df = pd.concat(curves, ignore_index=True) if curves else pd.DataFrame()
    return registry, raw, curve_df


def explode_groups(df):
    if df.empty:
        return df
    x = df.copy()
    x["experiment_group"] = x["experiment_groups"].fillna("").str.split(",")
    return x.explode("experiment_group")


def summarize(raw):
    done = raw[raw["status"] == "completed"].copy() if not raw.empty else raw
    if done.empty:
        return pd.DataFrame()
    done = explode_groups(done)
    keys = [
        "experiment_group", "task", "dataset", "method", "estimator", "M", "delta",
        "x_momentum", "y_momentum", "direction_distribution", "batch_size", "hidden",
        "c1", "c2", "lr_x", "lr_y"
    ]
    values = [
        "final_test_metric", "best_test_metric", "global_steps", "cum_train_forward_queries",
        "cum_train_backward_calls", "cum_train_sample_forwards", "train_wall_time_sec"
    ]
    agg = done.groupby(keys, dropna=False)[values].agg(["mean", "std", "count"]).reset_index()
    agg.columns = ["_".join(str(y) for y in x if str(y)) if isinstance(x, tuple) else str(x) for x in agg.columns]
    return agg.sort_values(["experiment_group", "task", "dataset"])


def status_dashboard(registry):
    return (registry.groupby(["phase", "assigned_gpu", "status"], dropna=False)
            .size().rename("runs").reset_index().sort_values(["phase", "assigned_gpu", "status"]))


def momentum_label(row):
    x = bool(row.get("x_momentum", True)); y = bool(row.get("y_momentum", True))
    return "full" if x and y else "no_x" if (not x and y) else "no_y" if (x and not y) else "no_both"


def curve_label(group, row):
    if group == "main_comparison": return str(row["method"])
    if group == "M_sweep": return f"M={int(row['M'])}"
    if group == "delta_sweep": return f"delta={row['delta']:g}"
    if group == "momentum": return momentum_label(row)
    if group == "perturbation_form": return str(row["estimator"])
    if group == "direction_ablation": return str(row["direction_distribution"])
    if group == "dimension": return f"hidden={int(row['hidden'])}"
    return str(row.get("method", "curve"))


def make_plots(curves):
    if curves.empty:
        return
    PLOTS.mkdir(parents=True, exist_ok=True)
    groups = ["main_comparison", "M_sweep", "delta_sweep", "momentum", "perturbation_form", "direction_ablation", "dimension"]
    x = explode_groups(curves)
    for group in groups:
        sub = x[x["experiment_group"] == group].copy()
        if sub.empty: continue
        sub["label"] = sub.apply(lambda r: curve_label(group, r), axis=1)
        for (task, dataset), ds in sub.groupby(["task", "dataset"]):
            plt.figure(figsize=(7.2, 4.8))
            for label, c in ds.groupby("label"):
                g = c.groupby("target_query")["test_metric"].agg(["mean", "std"]).reset_index().sort_values("target_query")
                xx = g["target_query"].to_numpy(); yy = g["mean"].to_numpy(); ss = g["std"].fillna(0).to_numpy()
                plt.plot(xx, yy, label=label, linewidth=1.8)
                plt.fill_between(xx, yy-ss, yy+ss, alpha=0.16)
            plt.xlabel("Training forward queries")
            plt.ylabel("Test AUC" if task == "auc" else "Test accuracy")
            plt.title(f"{group}: {task}/{dataset}")
            plt.grid(alpha=0.25); plt.legend(fontsize=8); plt.tight_layout()
            out = PLOTS / group
            out.mkdir(parents=True, exist_ok=True)
            plt.savefig(out / f"{task}_{dataset}.png", dpi=180)
            plt.close()


def write_excel(registry, raw, curves, summary):
    WORKBOOK.parent.mkdir(parents=True, exist_ok=True)
    group_sheets = {
        "Main Comparison": "main_comparison", "M Sweep": "M_sweep", "Delta Sweep": "delta_sweep",
        "Momentum": "momentum", "Perturbation": "perturbation_form", "Direction": "direction_ablation",
        "Coeff Batch": "coefficient_batch", "Dimension": "dimension",
        "M Delta": "M_delta_interaction", "M Momentum": "M_momentum_interaction", "FO LR Sweep": "fo_lr_sweep",
    }
    with pd.ExcelWriter(WORKBOOK, engine="openpyxl") as writer:
        pd.DataFrame({
            "说明": [
                "本工作簿由 rebuttal/aggregate_results.py 自动生成，请不要手工改 Raw Results。",
                "Run Registry：完整实验表；Status Dashboard：阶段/GPU/状态统计。",
                "Raw Results：每个 run 的最终结果；Raw Curves：每个 query checkpoint 的原始曲线。",
                "Summary 和各实验分页：三 seed 的 mean/std/count 汇总。",
                "主公平口径为 training forward queries；同时保留 backward、sample-forward 和 wall-clock。",
                "运行 status：pending/running/completed/failed。执行 refresh_results.sh 可随时刷新。",
            ]
        }).to_excel(writer, sheet_name="README", index=False)
        status_dashboard(registry).to_excel(writer, sheet_name="Status Dashboard", index=False)
        registry.to_excel(writer, sheet_name="Run Registry", index=False)
        raw.to_excel(writer, sheet_name="Raw Results", index=False)
        (curves if not curves.empty else pd.DataFrame(columns=["no_results_yet"])).to_excel(writer, sheet_name="Raw Curves", index=False)
        (summary if not summary.empty else pd.DataFrame(columns=["no_completed_results_yet"])).to_excel(writer, sheet_name="Summary", index=False)
        for sheet, group in group_sheets.items():
            part = summary[summary["experiment_group"] == group] if not summary.empty else pd.DataFrame()
            (part if not part.empty else pd.DataFrame(columns=["no_completed_results_yet"])).to_excel(writer, sheet_name=sheet[:31], index=False)
    wb = load_workbook(WORKBOOK)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    status_colors = {"completed": "C6EFCE", "running": "FFF2CC", "failed": "FFC7CE", "pending": "D9EAF7"}
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        if ws.max_row >= 1 and ws.max_column >= 1:
            ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = header_fill; cell.font = header_font; cell.alignment = Alignment(horizontal="center")
        for col in ws.columns:
            vals = [str(c.value) if c.value is not None else "" for c in list(col)[:200]]
            width = min(42, max(10, max((len(v) for v in vals), default=10) + 2))
            ws.column_dimensions[col[0].column_letter].width = width
        headers = {c.value: c.column for c in ws[1]}
        if "status" in headers:
            col = headers["status"]
            for row in range(2, ws.max_row + 1):
                cell = ws.cell(row, col); color = status_colors.get(str(cell.value))
                if color: cell.fill = PatternFill("solid", fgColor=color)
    wb.save(WORKBOOK)


def refresh():
    registry, raw, curves = collect()
    summary = summarize(raw)
    write_excel(registry, raw, curves, summary)
    make_plots(curves)
    counts = registry["status"].value_counts().to_dict()
    print(time.strftime("%F %T"), counts, "->", WORKBOOK, flush=True)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--watch", action="store_true")
    ap.add_argument("--interval", type=int, default=600)
    args = ap.parse_args()
    while True:
        try: refresh()
        except Exception as exc: print("refresh failed:", repr(exc), flush=True)
        if not args.watch: break
        time.sleep(args.interval)

if __name__ == "__main__": main()
