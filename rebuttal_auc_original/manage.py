#!/usr/bin/env python3
import argparse
import csv
import glob
import hashlib
import json
import math
import os
from pathlib import Path
import statistics
import subprocess
import sys
import time
from datetime import datetime

ROOT = Path(__file__).resolve().parents[1]
HERE = Path(__file__).resolve().parent
AUC_DIR = ROOT / "26.1.29.ZOFOGDA_AUC"
RUNS_DIR = HERE / "runs"
QUEUES_DIR = HERE / "queues"
LOGS_DIR = HERE / "logs"
MANIFEST = HERE / "manifest.json"
REGISTRY = HERE / "experiment_registry.csv"
WORKBOOK = HERE / "AUC_rebuttal_results.xlsx"
DATASETS = ["climate", "german", "svmguide1", "spambase", "a9a", "ijcnn1"]
SEEDS = [14, 42, 81]
EPOCHS = 500
BATCH_SIZE = 128
ZO_STEP = {
    "climate": (0.01, 0.01),
    "german": (0.005, 0.005),
    "svmguide1": (0.001, 0.001),
    "spambase": (0.04, 0.04),
    "a9a": (0.01, 0.01),
    "ijcnn1": (0.01, 0.01),
}
TTGDA_LR = {d: (0.01, 0.0001) for d in DATASETS}
TTGDA_LR["ijcnn1"] = (0.08, 0.0008)
TTSGDA_LR = {d: (0.01, 0.0001) for d in DATASETS}
METHOD_ORDER = {"TTGDA": 0, "TTSGDA": 1, "ZOFOSGDA": 2}


def now_text():
    return datetime.now().astimezone().isoformat(timespec="seconds")


def atomic_json(path, obj):
    path = Path(path)
    tmp = path.with_suffix(path.suffix + ".tmp")
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(obj, f, ensure_ascii=False, indent=2)
    os.replace(tmp, path)


def safe_value(v):
    return str(v).replace("-", "m").replace(".", "p")


def run_id(method, dataset, experiment, value, seed):
    base = f"auc_{method.lower()}_{dataset}_{experiment}_{safe_value(value)}_s{seed}"
    digest = hashlib.sha1(base.encode()).hexdigest()[:7]
    return f"{base}_{digest}"


def dataset_sizes():
    from sklearn.datasets import load_svmlight_file
    sizes = {}
    for d in DATASETS:
        train = AUC_DIR / "data" / d
        test = AUC_DIR / "data" / f"{d}.t"
        xtr, _ = load_svmlight_file(train)
        xte, _ = load_svmlight_file(test)
        sizes[d] = {"train": int(xtr.shape[0]), "test": int(xte.shape[0])}
    return sizes


def make_run(method, dataset, experiment, value, seed, sizes, **kwargs):
    q = int(kwargs.get("q", 5))
    n = sizes[dataset]["train"]
    multiplier = 2 * q + 1 if method == "ZOFOSGDA" else 1
    estimated_cost = EPOCHS * n * multiplier
    return {
        "run_id": run_id(method, dataset, experiment, value, seed),
        "method": method,
        "dataset": dataset,
        "experiment": experiment,
        "value": str(value),
        "seed": seed,
        "epochs": EPOCHS,
        "batch_size": BATCH_SIZE,
        "hidden": 128,
        "q": q,
        "delta": float(kwargs.get("delta", 0.001)),
        "x_momentum": bool(kwargs.get("x_momentum", True)),
        "y_momentum": bool(kwargs.get("y_momentum", True)),
        "lr_x": kwargs.get("lr_x"),
        "lr_y": kwargs.get("lr_y"),
        "gamma": kwargs.get("gamma"),
        "lam": kwargs.get("lam"),
        "n_train": sizes[dataset]["train"],
        "n_test": sizes[dataset]["test"],
        "estimated_cost": estimated_cost,
    }


def build(force=False):
    if MANIFEST.exists() and not force:
        runs = json.loads(MANIFEST.read_text(encoding="utf-8"))
        print(f"Manifest already exists: {len(runs)} runs")
        return runs
    sizes = dataset_sizes()
    runs = []
    for d in DATASETS:
        gx, gy = ZO_STEP[d]
        tx, ty = TTGDA_LR[d]
        sx, sy = TTSGDA_LR[d]
        for seed in SEEDS:
            runs.append(make_run("TTGDA", d, "baseline", "paper", seed, sizes, lr_x=tx, lr_y=ty))
            runs.append(make_run("TTSGDA", d, "baseline", "paper", seed, sizes, lr_x=sx, lr_y=sy))
            runs.append(make_run("ZOFOSGDA", d, "baseline", "paper", seed, sizes,
                                 q=5, delta=0.001, gamma=gx, lam=gy))
            for q in [1, 10, 15]:
                runs.append(make_run("ZOFOSGDA", d, "M", q, seed, sizes,
                                     q=q, delta=0.001, gamma=gx, lam=gy))
            for delta in [0.0001, 0.01]:
                runs.append(make_run("ZOFOSGDA", d, "delta", delta, seed, sizes,
                                     q=5, delta=delta, gamma=gx, lam=gy))
            for label, xm, ym in [("no_x", False, True), ("no_y", True, False),
                                  ("no_both", False, False)]:
                runs.append(make_run("ZOFOSGDA", d, "momentum", label, seed, sizes,
                                     q=5, delta=0.001, x_momentum=xm,
                                     y_momentum=ym, gamma=gx, lam=gy))
    assert len(runs) == 198, len(runs)
    assert len({r["run_id"] for r in runs}) == 198

    # Greedy cost balancing, with expensive jobs first.
    queues = [[] for _ in range(8)]
    loads = [0] * 8
    for r in sorted(runs, key=lambda x: x["estimated_cost"], reverse=True):
        gpu = min(range(8), key=lambda i: loads[i])
        r["gpu"] = gpu
        queues[gpu].append(r)
        loads[gpu] += r["estimated_cost"]

    HERE.mkdir(parents=True, exist_ok=True)
    RUNS_DIR.mkdir(parents=True, exist_ok=True)
    QUEUES_DIR.mkdir(parents=True, exist_ok=True)
    LOGS_DIR.mkdir(parents=True, exist_ok=True)
    atomic_json(MANIFEST, runs)
    for gpu, queue in enumerate(queues):
        atomic_json(QUEUES_DIR / f"gpu{gpu}.json", queue)
    fields = ["run_id", "gpu", "dataset", "method", "experiment", "value", "seed",
              "epochs", "batch_size", "q", "delta", "x_momentum", "y_momentum",
              "lr_x", "lr_y", "gamma", "lam", "n_train", "n_test", "estimated_cost"]
    with open(REGISTRY, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=fields, lineterminator="\n")
        w.writeheader()
        for r in runs:
            w.writerow({k: r.get(k) for k in fields})
    print(f"Built {len(runs)} unique runs; queue sizes={[len(q) for q in queues]}")
    print(f"Relative loads={[round(x / min(loads), 3) for x in loads]}")
    return runs


def command_for(r, gpu, out_dir):
    common = [
        sys.executable,
        str(AUC_DIR / ({"TTGDA": "ldx_ttgda_auc.py", "TTSGDA": "ldx_ttsgda_auc.py",
                        "ZOFOSGDA": "ldx_zofogda_auc.py"}[r["method"]])),
        "--dataset", r["dataset"], "--data_dir", "./data", "--hidden", str(r["hidden"]),
        "--epochs", str(r["epochs"]), "--batch_size", str(r["batch_size"]),
        "--standardize", "--seed", str(r["seed"]), "--device", f"cuda:{gpu}",
        "--log_dir", str(out_dir),
    ]
    if r["method"] == "ZOFOSGDA":
        common += [
            "--log_prefix", "zofoGDA_auc", "--eta_k", "10", "--eta_m", "1000",
            "--c1", "3", "--c2", "3", "--gamma", str(r["gamma"]),
            "--lam", str(r["lam"]), "--zo_delta", str(r["delta"]), "--zo_q", str(r["q"]),
        ]
        if not r["x_momentum"]:
            common.append("--no_x_momentum")
        if not r["y_momentum"]:
            common.append("--no_y_momentum")
    else:
        common += ["--log_prefix", r["method"].lower() + "_auc",
                   "--lr_x", str(r["lr_x"]), "--lr_y", str(r["lr_y"]),
                   "--y_steps_per_batch", "1", "--v_clip", "10"]
    return common


def numeric_rows(csv_path):
    rows = []
    with open(csv_path, newline="") as f:
        for row in csv.DictReader(f):
            if (row.get("epoch") or "").isdigit():
                rows.append(row)
    return rows


def query_metrics(r):
    n = r["n_train"]
    epochs = r["epochs"]
    if r["method"] == "TTGDA":
        steps = epochs
        calls = steps
        sample_queries = epochs * n
        backwards = steps
    elif r["method"] == "TTSGDA":
        steps = epochs * math.ceil(n / r["batch_size"])
        calls = steps
        sample_queries = epochs * n
        backwards = steps
    else:
        steps = epochs * math.ceil(n / r["batch_size"])
        calls = steps * (2 * r["q"] + 1)
        sample_queries = epochs * n * (2 * r["q"] + 1)
        backwards = steps
    return {"optimizer_steps": steps, "train_forward_calls": calls,
            "train_backward_calls": backwards, "train_sample_function_queries": sample_queries}


def gpu_context_count(gpu):
    result = subprocess.run(
        ["nvidia-smi", "-i", str(gpu), "--query-compute-apps=pid",
         "--format=csv,noheader,nounits"],
        text=True, capture_output=True,
    )
    if result.returncode != 0:
        raise RuntimeError(f"nvidia-smi failed for GPU {gpu}: {result.stderr.strip()}")
    return len([line for line in result.stdout.splitlines() if line.strip()])


def wait_for_empty_gpu(gpu, run_id):
    announced = False
    empty_checks = 0
    while empty_checks < 3:
        count = gpu_context_count(gpu)
        if count == 0:
            empty_checks += 1
        else:
            empty_checks = 0
            if not announced:
                print(f"[{now_text()}] GPU {gpu} busy ({count} contexts); {run_id} waits", flush=True)
                announced = True
        time.sleep(1.0)


def worker(gpu):
    queue_path = QUEUES_DIR / f"gpu{gpu}.json"
    if not queue_path.exists():
        build()
    queue = json.loads(queue_path.read_text(encoding="utf-8"))
    print(f"[{now_text()}] GPU {gpu}: {len(queue)} planned runs", flush=True)
    for pos, r in enumerate(queue, 1):
        out_dir = RUNS_DIR / r["run_id"]
        out_dir.mkdir(parents=True, exist_ok=True)
        summary_path = out_dir / "summary.json"
        if summary_path.exists():
            try:
                old = json.loads(summary_path.read_text(encoding="utf-8"))
                if (old.get("status") == "completed" and
                        old.get("epochs_completed") == r["epochs"] and
                        old.get("gpu_exclusive") is True and
                        Path(old.get("epoch_times_path", "")).exists()):
                    print(f"[{pos}/{len(queue)}] skip {r['run_id']}", flush=True)
                    continue
            except Exception:
                pass
        atomic_json(out_dir / "config.json", r)
        cmd = command_for(r, gpu, out_dir)
        log_path = LOGS_DIR / f"{r['run_id']}.log"
        attempt = 0
        while True:
            attempt += 1
            wait_for_empty_gpu(gpu, r["run_id"])
            started = time.perf_counter()
            started_at = now_text()
            atomic_json(out_dir / "running.json", {
                "status": "running", "gpu": gpu, "started_at": started_at,
                "run_id": r["run_id"], "exclusive_attempt": attempt,
            })
            print(f"[{pos}/{len(queue)}] start {r['run_id']} exclusive_attempt={attempt}", flush=True)
            audit_path = out_dir / f"gpu_audit_attempt{attempt}.csv"
            epoch_times_path = out_dir / "epoch_times.csv"
            contaminated = False
            max_contexts = 0
            last_audit = 0.0
            last_size = -1
            seen_epochs = 0
            with open(log_path, "w", encoding="utf-8") as log, \
                    open(audit_path, "w", newline="", encoding="utf-8") as audit, \
                    open(epoch_times_path, "w", newline="", encoding="utf-8") as et:
                log.write("COMMAND: " + " ".join(cmd) + "\n")
                log.flush()
                aw = csv.writer(audit)
                aw.writerow(["elapsed_wall_time_sec", "gpu_context_count", "exclusive"])
                ew = csv.writer(et)
                ew.writerow(["epoch", "elapsed_wall_time_sec", "time_ms", "observed_at"])
                proc = subprocess.Popen(cmd, cwd=AUC_DIR, stdout=log, stderr=subprocess.STDOUT)
                while proc.poll() is None:
                    elapsed = time.perf_counter() - started
                    csvs = list(out_dir.glob("*.csv"))
                    train_csvs = [x for x in csvs if x.name != "epoch_times.csv" and not x.name.startswith("gpu_audit")]
                    if len(train_csvs) == 1:
                        size = train_csvs[0].stat().st_size
                        if size != last_size:
                            last_size = size
                            rows = numeric_rows(train_csvs[0])
                            if len(rows) > seen_epochs:
                                stamp = now_text()
                                for row in rows[seen_epochs:]:
                                    ew.writerow([int(row["epoch"]), elapsed, elapsed * 1000.0, stamp])
                                et.flush()
                                seen_epochs = len(rows)
                    if elapsed - last_audit >= 1.0:
                        count = gpu_context_count(gpu)
                        max_contexts = max(max_contexts, count)
                        aw.writerow([elapsed, count, count <= 1])
                        audit.flush()
                        last_audit = elapsed
                        if count > 1:
                            contaminated = True
                            log.write(f"\nGPU EXCLUSIVITY VIOLATION: {count} contexts; restarting run.\n")
                            log.flush()
                            proc.terminate()
                            try:
                                proc.wait(timeout=20)
                            except subprocess.TimeoutExpired:
                                proc.kill()
                                proc.wait()
                            break
                    time.sleep(0.02)
                returncode = proc.wait()
                duration = time.perf_counter() - started
                csvs = [x for x in out_dir.glob("*.csv") if x.name != "epoch_times.csv" and not x.name.startswith("gpu_audit")]
                if len(csvs) == 1:
                    rows_now = numeric_rows(csvs[0])
                    if len(rows_now) > seen_epochs:
                        stamp = now_text()
                        for row in rows_now[seen_epochs:]:
                            ew.writerow([int(row["epoch"]), duration, duration * 1000.0, stamp])
                        et.flush()
                        seen_epochs = len(rows_now)
            if contaminated:
                atomic_json(out_dir / "running.json", {
                    "status": "waiting_for_exclusive_restart", "gpu": gpu,
                    "run_id": r["run_id"], "exclusive_attempt": attempt,
                    "max_gpu_contexts": max_contexts,
                })
                print(f"[{pos}/{len(queue)}] contamination detected; restart {r['run_id']}", flush=True)
                continue
            try:
                if returncode != 0 or len(csvs) != 1:
                    raise RuntimeError(f"returncode={returncode}, csv_count={len(csvs)}")
                rows = numeric_rows(csvs[0])
                epoch_times = numeric_rows(epoch_times_path)
                if len(rows) != r["epochs"]:
                    raise RuntimeError(f"expected {r['epochs']} epochs, got {len(rows)}")
                if len(epoch_times) != r["epochs"]:
                    raise RuntimeError(f"expected {r['epochs']} epoch timings, got {len(epoch_times)}")
                aucs = [float(x["test_auc"]) for x in rows]
                losses = [float(x["train_loss"]) for x in rows]
                summary = {
                    **r, "status": "completed", "started_at": started_at,
                    "finished_at": now_text(), "total_wall_time_sec": duration,
                    "epochs_completed": len(rows), "final_test_auc": aucs[-1],
                    "best_test_auc": max(aucs), "best_epoch": aucs.index(max(aucs)) + 1,
                    "final_train_loss": losses[-1], "csv_path": str(csvs[0]),
                    "epoch_times_path": str(epoch_times_path), "gpu_audit_path": str(audit_path),
                    "gpu_exclusive": True, "max_gpu_contexts": max_contexts,
                    "exclusive_attempts": attempt, "log_path": str(log_path), **query_metrics(r),
                }
                atomic_json(summary_path, summary)
                print(f"[{pos}/{len(queue)}] done {r['run_id']} AUC={aucs[-1]:.6f} time={duration:.1f}s", flush=True)
            except Exception as exc:
                atomic_json(summary_path, {**r, "status": "failed", "started_at": started_at,
                                          "finished_at": now_text(), "total_wall_time_sec": duration,
                                          "gpu_exclusive": max_contexts <= 1,
                                          "max_gpu_contexts": max_contexts,
                                          "error": repr(exc), "log_path": str(log_path)})
                print(f"[{pos}/{len(queue)}] FAILED {r['run_id']}: {exc!r}", flush=True)
            break
        try:
            (out_dir / "running.json").unlink()
        except FileNotFoundError:
            pass
    print(f"[{now_text()}] GPU {gpu}: queue finished", flush=True)

def old_baseline(dataset, method, seed):
    prefix = {"TTGDA": "ttgda_auc_", "TTSGDA": "ttsgda_auc_", "ZOFOSGDA": "zofoGDA_auc_"}[method]
    matches = glob.glob(str(AUC_DIR / "result_plot" / dataset / f"{prefix}{dataset}_epochs500_seed{seed}_*.csv"))
    if len(matches) != 1:
        return None
    rows = numeric_rows(matches[0])
    aucs = [float(x["test_auc"]) for x in rows]
    return {"status": "paper_existing", "final_test_auc": aucs[-1], "best_test_auc": max(aucs),
            "best_epoch": aucs.index(max(aucs)) + 1, "epochs_completed": len(rows),
            "total_wall_time_sec": None, "train_sample_function_queries": None,
            "csv_path": matches[0]}


def load_state(r):
    p = RUNS_DIR / r["run_id"] / "summary.json"
    if p.exists():
        try:
            return json.loads(p.read_text(encoding="utf-8"))
        except Exception:
            return {"status": "invalid_summary"}
    running = RUNS_DIR / r["run_id"] / "running.json"
    return {"status": "running" if running.exists() else "pending"}


def refresh():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.formatting.rule import ColorScaleRule

    runs = build()
    by_key = {(r["dataset"], r["method"], r["experiment"], r["value"], r["seed"]): r for r in runs}
    states = {r["run_id"]: load_state(r) for r in runs}

    def get_run(dataset, method, experiment, value, seed):
        if experiment == "M" and str(value) == "5":
            experiment, value = "baseline", "paper"
        elif experiment == "delta" and str(value) in {"0.001", "1e-3"}:
            experiment, value = "baseline", "paper"
        elif experiment == "momentum" and value == "full":
            experiment, value = "baseline", "paper"
        return by_key[(dataset, method, experiment, str(value), seed)]

    def result_for(r, allow_paper=False):
        # Formal outputs only use independently rerun, auditable results.
        s = states[r["run_id"]]
        if s.get("status") == "completed" and s.get("gpu_exclusive") is True:
            return s, "本次干净复跑"
        return s, s.get("status", "pending")

    wb = Workbook()
    wb.remove(wb.active)
    blue = "1F4E78"; light_blue = "D9EAF7"; green = "E2F0D9"; yellow = "FFF2CC"; red = "F4CCCC"
    thin = Side(style="thin", color="D9E1F2")

    def setup(ws, headers):
        ws.append(headers)
        for c in ws[1]:
            c.fill = PatternFill("solid", fgColor=blue)
            c.font = Font(color="FFFFFF", bold=True)
            c.alignment = Alignment(horizontal="center", vertical="center")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False
        ws.row_dimensions[1].height = 26

    def finish(ws, widths):
        for col, width in widths.items():
            ws.column_dimensions[col].width = width
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.border = Border(bottom=thin)
                c.alignment = Alignment(vertical="center", wrap_text=True)

    # User-requested wide main table. Each dataset has explicit mean and variance columns.
    main_specs = [
        ("TTGDA-原配置复跑", "TTGDA", "baseline", "paper"),
        ("TTSGDA-原配置复跑", "TTSGDA", "baseline", "paper"),
        ("ZOFOGDA-原配置复跑", "ZOFOSGDA", "baseline", "paper"),
        ("ZOFOGDA-M=1", "ZOFOSGDA", "M", 1),
        ("ZOFOGDA-M=10", "ZOFOSGDA", "M", 10),
        ("ZOFOGDA-M=15", "ZOFOSGDA", "M", 15),
        ("ZOFOGDA-delta=1e-4", "ZOFOSGDA", "delta", 0.0001),
        ("ZOFOGDA-delta=1e-2", "ZOFOSGDA", "delta", 0.01),
        ("ZOFOGDA-no-x", "ZOFOSGDA", "momentum", "no_x"),
        ("ZOFOGDA-no-y", "ZOFOSGDA", "momentum", "no_y"),
        ("ZOFOGDA-no-both", "ZOFOSGDA", "momentum", "no_both"),
    ]

    def condition_values(dataset, method, experiment, value):
        vals = []
        for seed in SEEDS:
            r = get_run(dataset, method, experiment, value, seed)
            result, _ = result_for(r)
            if result.get("final_test_auc") is not None:
                vals.append(float(result["final_test_auc"]))
        return vals

    ws = wb.create_sheet("最终主表", 0)
    ws.merge_cells("A1:M1")
    ws["A1"] = "AUC"
    ws["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=blue)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells("A2:A3")
    ws["A2"] = "方法"
    for j, d in enumerate(DATASETS):
        c = 2 + 2 * j
        ws.merge_cells(start_row=2, start_column=c, end_row=2, end_column=c + 1)
        ws.cell(2, c, d.capitalize())
        ws.cell(3, c, "均值")
        ws.cell(3, c + 1, "方差")
    for row in ws.iter_rows(min_row=2, max_row=3):
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=light_blue)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
    for label, method, experiment, value in main_specs:
        row = [label]
        for d in DATASETS:
            vals = condition_values(d, method, experiment, value)
            row.extend([statistics.mean(vals) if len(vals) == 3 else "",
                        statistics.pvariance(vals) if len(vals) == 3 else ""])
        ws.append(row)
    ws.freeze_panes = "B4"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 27
    for col in range(2, 14):
        ws.column_dimensions[chr(64 + col)].width = 13
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row, min_col=2, max_col=13):
        for cell in row:
            cell.number_format = "0.000000"
            cell.alignment = Alignment(horizontal="center")
    ws["A16"] = "说明：方差使用3个seed的总体方差（ddof=0）；论文常见的 mean±std 可直接使用下一张表。"
    ws.merge_cells("A16:M16")
    ws["A16"].font = Font(italic=True, color="666666")

    ws = wb.create_sheet("主表_mean±std", 1)
    setup(ws, ["方法"] + [d.capitalize() for d in DATASETS])
    for label, method, experiment, value in main_specs:
        row = [label]
        for d in DATASETS:
            vals = condition_values(d, method, experiment, value)
            row.append(f"{statistics.mean(vals):.4f} ± {statistics.pstdev(vals):.4f}" if len(vals) == 3 else "")
        ws.append(row)
    finish(ws, {"A":27, "B":21, "C":21, "D":21, "E":21, "F":21, "G":21})

    # Sheet 1: baseline, always grouped in three consecutive method rows.
    ws = wb.create_sheet("三方法对比")
    setup(ws, ["数据集", "方法", "消融变量", "对应值", "所做任务", "结果（Final AUC）",
               "Best AUC", "完成种子", "平均运行时间(min)", "平均Sample Queries", "数据来源"])
    for d in DATASETS:
        for method in ["TTGDA", "TTSGDA", "ZOFOSGDA"]:
            vals=[]; best=[]; times=[]; queries=[]; sources=[]
            for seed in SEEDS:
                r=get_run(d,method,"baseline","paper",seed); s,source=result_for(r)
                if s.get("final_test_auc") is not None:
                    vals.append(float(s["final_test_auc"])); best.append(float(s["best_test_auc"])); sources.append(source)
                if s.get("total_wall_time_sec") is not None: times.append(float(s["total_wall_time_sec"])/60)
                if s.get("train_sample_function_queries") is not None: queries.append(int(s["train_sample_function_queries"]))
            result = f"{statistics.mean(vals):.4f} ± {statistics.pstdev(vals):.4f}" if vals else ""
            best_text = f"{statistics.mean(best):.4f} ± {statistics.pstdev(best):.4f}" if best else ""
            ws.append([d,method,"无（原配置复跑）","原配置","AUC最大化",result,best_text,
                       f"{len(vals)}/3", round(statistics.mean(times),2) if times else "",
                       round(statistics.mean(queries)) if queries else "", "/".join(sorted(set(sources)))])
        for cell in ws[ws.max_row-2]: cell.fill=PatternFill("solid",fgColor=light_blue)
    finish(ws,{"A":15,"B":14,"C":22,"D":15,"E":16,"F":22,"G":22,"H":12,"I":20,"J":22,"K":18})

    # Sheet 2: ZO ablations, concise and complete.
    ws = wb.create_sheet("消融结果")
    setup(ws,["数据集","方法","消融变量","对应值","所做任务","结果（Final AUC）","Best AUC",
              "完成种子","平均运行时间(min)","平均Sample Queries","状态/来源"])
    conditions = [("M",v) for v in [1,5,10,15]] + [("delta",v) for v in [0.0001,0.001,0.01]] + [("momentum",v) for v in ["full","no_x","no_y","no_both"]]
    for d in DATASETS:
        for exp,val in conditions:
            vals=[]; best=[]; times=[]; queries=[]; sources=[]
            for seed in SEEDS:
                r=get_run(d,"ZOFOSGDA",exp,val,seed); s,source=result_for(r)
                if s.get("final_test_auc") is not None:
                    vals.append(float(s["final_test_auc"])); best.append(float(s["best_test_auc"])); sources.append(source)
                if s.get("total_wall_time_sec") is not None: times.append(float(s["total_wall_time_sec"])/60)
                if s.get("train_sample_function_queries") is not None: queries.append(int(s["train_sample_function_queries"]))
            result=f"{statistics.mean(vals):.4f} ± {statistics.pstdev(vals):.4f}" if vals else ""
            best_text=f"{statistics.mean(best):.4f} ± {statistics.pstdev(best):.4f}" if best else ""
            ws.append([d,"ZOFOSGDA",exp,str(val),"AUC最大化",result,best_text,f"{len(vals)}/3",
                       round(statistics.mean(times),2) if times else "",round(statistics.mean(queries)) if queries else "",
                       "/".join(sorted(set(sources))) if sources else "等待运行"])
    finish(ws,{"A":15,"B":14,"C":15,"D":15,"E":16,"F":22,"G":22,"H":12,"I":20,"J":22,"K":18})

    # Sheet 3: every unique run for auditing and runtime.
    ws = wb.create_sheet("每次运行")
    headers=["Run ID","GPU","数据集","方法","实验","值","Seed","状态","GPU独占","最大GPU上下文","独占尝试次数",
             "Final AUC","Best AUC","Best Epoch","运行时间(s)","Epochs","优化步数","Forward Calls","Backward Calls",
             "Sample Queries","结果CSV","逐Epoch时间","GPU审计","日志"]
    setup(ws,headers)
    status_counts={}
    for r in sorted(runs,key=lambda x:(DATASETS.index(x["dataset"]),METHOD_ORDER[x["method"]],x["experiment"],x["value"],x["seed"])):
        s=states[r["run_id"]]; status=s.get("status","pending"); status_counts[status]=status_counts.get(status,0)+1
        ws.append([r["run_id"],r["gpu"],r["dataset"],r["method"],r["experiment"],r["value"],r["seed"],status,
                   s.get("gpu_exclusive",""),s.get("max_gpu_contexts",""),s.get("exclusive_attempts",""),
                   s.get("final_test_auc",""),s.get("best_test_auc",""),s.get("best_epoch",""),
                   round(s.get("total_wall_time_sec",0),3) if s.get("total_wall_time_sec") is not None else "",
                   s.get("epochs_completed",0),s.get("optimizer_steps",""),s.get("train_forward_calls",""),
                   s.get("train_backward_calls",""),s.get("train_sample_function_queries",""),
                   s.get("csv_path",""),s.get("epoch_times_path",""),s.get("gpu_audit_path",""),s.get("log_path","")])
        fill={"completed":green,"running":yellow,"failed":red}.get(status)
        if fill: ws.cell(ws.max_row,8).fill=PatternFill("solid",fgColor=fill)
    finish(ws,{"A":52,"B":7,"C":14,"D":14,"E":12,"F":12,"G":8,"H":12,"I":14,"J":14,"K":12,
               "L":16,"M":10,"N":15,"O":16,"P":16,"Q":20,"R":55,"S":55})
    ws.conditional_formatting.add(f"I2:I{ws.max_row}",ColorScaleRule(start_type="min",start_color="F8696B",mid_type="percentile",mid_value=50,mid_color="FFEB84",end_type="max",end_color="63BE7B"))

    ws = wb.create_sheet("进度")
    setup(ws,["项目","数量/内容"])
    completed=status_counts.get("completed",0); running=status_counts.get("running",0); failed=status_counts.get("failed",0)
    rows=[("更新时间",now_text()),("计划唯一Runs",len(runs)),("已完成",completed),("运行中",running),
          ("失败",failed),("等待",len(runs)-completed-running-failed),("完成比例",f"{completed/len(runs):.1%}"),
          ("实验范围","仅AUC：6数据集，500 epochs，seeds 14/42/81"),
          ("原始训练入口","ldx_zofogda_auc.py / ldx_ttgda_auc.py / ldx_ttsgda_auc.py"),
          ("说明","正式表仅使用本次独占GPU、seeds 14/42/81 的可审计复跑；历史精选结果不进入汇总。"),
          ("历史数据说明","旧result_plot仅保留作历史审计，不作为正式结果或默认绘图数据源。")]
    for row in rows: ws.append(row)
    finish(ws,{"A":24,"B":80})

    tmp=WORKBOOK.with_suffix(".tmp.xlsx")
    wb.save(tmp); os.replace(tmp,WORKBOOK)
    print(f"{now_text()} status={status_counts} -> {WORKBOOK}")
    return status_counts


def session_exists(name):
    return subprocess.run(["tmux","has-session","-t",name],stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL).returncode==0


def launch():
    build(); refresh()
    for gpu in range(8):
        name=f"auc_original_gpu{gpu}"
        if not session_exists(name):
            cmd=f"cd {ROOT} && {sys.executable} {Path(__file__).resolve()} worker --gpu {gpu}"
            subprocess.run(["tmux","new-session","-d","-s",name,cmd],check=True)
    report="auc_original_report"
    if not session_exists(report):
        cmd=f"cd {ROOT} && {sys.executable} {Path(__file__).resolve()} reporter"
        subprocess.run(["tmux","new-session","-d","-s",report,cmd],check=True)
    print("Launched tmux sessions auc_original_gpu0..7 and auc_original_report")


def reporter():
    while True:
        status=refresh()
        terminal=status.get("completed",0)+status.get("failed",0)
        if terminal>=198:
            break
        time.sleep(120)


def status():
    runs=build(); counts={}
    for r in runs:
        s=load_state(r).get("status","pending"); counts[s]=counts.get(s,0)+1
    print(json.dumps(counts,ensure_ascii=False))
    for gpu in range(8):
        name=f"auc_original_gpu{gpu}"
        print(name,"alive" if session_exists(name) else "stopped")


def main():
    p=argparse.ArgumentParser()
    sub=p.add_subparsers(dest="cmd",required=True)
    sub.add_parser("build").add_argument("--force",action="store_true")
    w=sub.add_parser("worker"); w.add_argument("--gpu",type=int,required=True,choices=range(8))
    sub.add_parser("refresh"); sub.add_parser("reporter"); sub.add_parser("launch"); sub.add_parser("status")
    args=p.parse_args()
    if args.cmd=="build": build(args.force)
    elif args.cmd=="worker": worker(args.gpu)
    elif args.cmd=="refresh": refresh()
    elif args.cmd=="reporter": reporter()
    elif args.cmd=="launch": launch()
    elif args.cmd=="status": status()

if __name__=="__main__": main()
