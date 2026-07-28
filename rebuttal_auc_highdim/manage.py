#!/usr/bin/env python3
"""Manage the 60 high-dimensional ZOFOGDA M-ablation runs."""
from __future__ import annotations

import argparse
import csv
import hashlib
import importlib.util
import json
import math
import os
from pathlib import Path
import statistics
import subprocess
import sys
import time
from datetime import datetime


ROOT = Path(__file__).resolve().parents[1]
HERE = Path(__file__).resolve().parent
AUC_DIR = ROOT / "26.1.29.ZOFOGDA_AUC"
DATA_DIR = AUC_DIR / "data_highdim"
RUNS_DIR = HERE / "runs"
QUEUES_DIR = HERE / "queues"
LOGS_DIR = HERE / "logs"
MANIFEST = HERE / "manifest.json"
REGISTRY = HERE / "experiment_registry.csv"
WORKBOOK = HERE / "AUC_highdim_M_results.xlsx"

DATASETS = ["madelon", "gisette", "arcene", "dexter"]
FEATURE_DIMS = {"madelon": 500, "gisette": 5000, "arcene": 10000, "dexter": 20000}
EXPECTED_SIZES = {
    "madelon": (2000, 600),
    "gisette": (6000, 1000),
    "arcene": (100, 100),
    "dexter": (300, 300),
}
SQRT_D_M = {name: math.isqrt(dim) for name, dim in FEATURE_DIMS.items()}
STEP = {
    "madelon": (0.01, 0.01),
    "gisette": (0.003, 0.003),
    "arcene": (0.005, 0.005),
    "dexter": (0.003, 0.003),
}
SEEDS = [14, 42, 81]
EPOCHS = 500
BATCH_SIZE = 128
M_VALUES = [1, 5, 10, 15, "sqrt_d"]


def load_core():
    path = ROOT / "rebuttal_auc_original" / "manage.py"
    spec = importlib.util.spec_from_file_location("auc_original_core", path)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


core = load_core()
core.HERE = HERE
core.AUC_DIR = AUC_DIR
core.RUNS_DIR = RUNS_DIR
core.QUEUES_DIR = QUEUES_DIR
core.LOGS_DIR = LOGS_DIR
core.MANIFEST = MANIFEST
core.REGISTRY = REGISTRY
core.WORKBOOK = WORKBOOK
core.DATASETS = DATASETS
core.FEATURE_DIMS = FEATURE_DIMS
core.SQRT_D_M = SQRT_D_M
core.SEEDS = SEEDS
core.EPOCHS = EPOCHS
core.BATCH_SIZE = BATCH_SIZE


def now_text() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def atomic_json(path: Path, value) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(json.dumps(value, ensure_ascii=False, indent=2))
    os.replace(tmp, path)


def run_id(dataset: str, value, seed: int) -> str:
    token = f"auc_zofosgda_{dataset}_highdim_M_{value}_s{seed}"
    return token + "_" + hashlib.sha1(token.encode()).hexdigest()[:7]


def dataset_sizes() -> dict:
    from sklearn.datasets import load_svmlight_file

    sizes = {}
    for dataset in DATASETS:
        dim = FEATURE_DIMS[dataset]
        xtr, ytr = load_svmlight_file(DATA_DIR / dataset, n_features=dim)
        xte, yte = load_svmlight_file(DATA_DIR / f"{dataset}.t", n_features=dim)
        expected_train, expected_test = EXPECTED_SIZES[dataset]
        if xtr.shape != (expected_train, dim) or xte.shape != (expected_test, dim):
            raise RuntimeError(f"{dataset}: unexpected shapes {xtr.shape}/{xte.shape}")
        if set(ytr) != {-1.0, 1.0} or set(yte) != {-1.0, 1.0}:
            raise RuntimeError(f"{dataset}: labels are not binary +/-1")
        sizes[dataset] = {"train": expected_train, "test": expected_test, "features": dim}
    return sizes


def make_run(dataset: str, value, seed: int, sizes: dict) -> dict:
    q = SQRT_D_M[dataset] if value == "sqrt_d" else int(value)
    gamma, lam = STEP[dataset]
    n_train = sizes[dataset]["train"]
    dim = sizes[dataset]["features"]
    estimated_cost = EPOCHS * n_train * (2 * q + 1) * dim
    return {
        "run_id": run_id(dataset, value, seed),
        "method": "ZOFOSGDA",
        "dataset": dataset,
        "experiment": "highdim_M",
        "value": str(value),
        "seed": seed,
        "epochs": EPOCHS,
        "batch_size": BATCH_SIZE,
        "hidden": 128,
        "q": q,
        "delta": 0.001,
        "x_momentum": True,
        "y_momentum": True,
        "momentum_mode": "scheduled",
        "momentum_decay": 0.1,
        "lr_x": None,
        "lr_y": None,
        "gamma": gamma,
        "lam": lam,
        "n_train": n_train,
        "n_test": sizes[dataset]["test"],
        "feature_dim": dim,
        "estimated_cost": estimated_cost,
    }


def write_registry(runs: list[dict]) -> None:
    fields = [
        "run_id", "gpu", "dataset", "method", "experiment", "value", "seed",
        "epochs", "batch_size", "feature_dim", "q", "delta", "x_momentum",
        "y_momentum", "momentum_mode", "gamma", "lam", "n_train", "n_test",
        "estimated_cost",
    ]
    with REGISTRY.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fields, lineterminator="\n")
        writer.writeheader()
        for run in runs:
            writer.writerow({field: run.get(field) for field in fields})


def build(force: bool = False) -> list[dict]:
    if MANIFEST.exists() and not force:
        runs = json.loads(MANIFEST.read_text())
        if len(runs) != 60 or len({run["run_id"] for run in runs}) != 60:
            raise RuntimeError("Existing high-dimensional manifest is invalid")
        print(f"Manifest already exists: {len(runs)} runs")
        return runs

    sizes = dataset_sizes()
    runs = [
        make_run(dataset, value, seed, sizes)
        for dataset in DATASETS
        for value in M_VALUES
        for seed in SEEDS
    ]
    if len(runs) != 60 or len({run["run_id"] for run in runs}) != 60:
        raise RuntimeError("Expected 60 unique high-dimensional runs")

    queues = [[] for _ in range(8)]
    loads = [0] * 8
    for run in sorted(runs, key=lambda item: item["estimated_cost"], reverse=True):
        gpu = min(range(8), key=lambda index: loads[index])
        run["gpu"] = gpu
        queues[gpu].append(run)
        loads[gpu] += run["estimated_cost"]

    RUNS_DIR.mkdir(parents=True, exist_ok=True)
    QUEUES_DIR.mkdir(parents=True, exist_ok=True)
    LOGS_DIR.mkdir(parents=True, exist_ok=True)
    atomic_json(MANIFEST, runs)
    for gpu, queue in enumerate(queues):
        atomic_json(QUEUES_DIR / f"gpu{gpu}.json", queue)
    write_registry(runs)
    print(f"Built 60 runs; queue sizes={[len(queue) for queue in queues]}")
    print(f"Relative loads={[round(load / min(loads), 3) for load in loads]}")
    return runs


original_command_for = core.command_for


def command_for(run: dict, gpu: int, out_dir: Path) -> list[str]:
    command = original_command_for(run, gpu, out_dir)
    data_index = command.index("--data_dir") + 1
    command[data_index] = "./data_highdim"
    return command


core.command_for = command_for
core.build = build


def load_state(run: dict) -> dict:
    return core.load_state(run)


def refresh() -> dict:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    runs = build()
    states = {run["run_id"]: load_state(run) for run in runs}
    lookup = {
        (run["dataset"], run["value"], run["seed"]): run
        for run in runs
    }
    workbook = Workbook()
    main = workbook.active
    main.title = "主表_mean±std"
    main.append(["方法"] + [dataset.capitalize() for dataset in DATASETS])
    labels = [
        ("ZOFOGDA-原配置(M=5)", "5"),
        ("ZOFOGDA-M=1", "1"),
        ("ZOFOGDA-M=10", "10"),
        ("ZOFOGDA-M=15", "15"),
        ("ZOFOGDA-M=floor(sqrt(d))", "sqrt_d"),
    ]
    for label, value in labels:
        row = [label]
        for dataset in DATASETS:
            values = []
            for seed in SEEDS:
                state = states[lookup[(dataset, value, seed)]["run_id"]]
                if state.get("status") == "completed" and state.get("gpu_exclusive") is True:
                    values.append(float(state["final_test_auc"]))
            row.append(
                f"{statistics.mean(values):.4f} ± {statistics.pstdev(values):.4f}"
                if len(values) == 3 else ""
            )
        main.append(row)

    blue = "1F4E78"
    for cell in main[1]:
        cell.fill = PatternFill("solid", fgColor=blue)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    main.freeze_panes = "B2"
    main.column_dimensions["A"].width = 31
    for column in ["B", "C", "D", "E"]:
        main.column_dimensions[column].width = 22

    detail = workbook.create_sheet("每次运行")
    detail.append([
        "Run ID", "GPU", "数据集", "M标签", "实际M", "Seed", "状态",
        "Final AUC", "Best AUC", "运行时间(s)", "GPU独占", "Epochs",
        "gamma", "lambda", "结果CSV", "逐Epoch时间", "GPU审计", "日志",
    ])
    counts = {}
    for run in sorted(runs, key=lambda item: (DATASETS.index(item["dataset"]), item["q"], item["seed"])):
        state = states[run["run_id"]]
        status = state.get("status", "pending")
        counts[status] = counts.get(status, 0) + 1
        detail.append([
            run["run_id"], run["gpu"], run["dataset"], run["value"], run["q"],
            run["seed"], status, state.get("final_test_auc", ""),
            state.get("best_test_auc", ""), state.get("total_wall_time_sec", ""),
            state.get("gpu_exclusive", ""), state.get("epochs_completed", ""),
            run["gamma"], run["lam"], state.get("csv_path", ""),
            state.get("epoch_times_path", ""), state.get("gpu_audit_path", ""),
            state.get("log_path", ""),
        ])
    for cell in detail[1]:
        cell.fill = PatternFill("solid", fgColor=blue)
        cell.font = Font(color="FFFFFF", bold=True)
    detail.freeze_panes = "A2"
    detail.auto_filter.ref = detail.dimensions

    progress = workbook.create_sheet("配置与进度")
    completed = counts.get("completed", 0)
    running = counts.get("running", 0)
    failed = counts.get("failed", 0)
    rows = [
        ("更新时间", now_text()),
        ("计划Runs", 60),
        ("已完成", completed),
        ("运行中", running),
        ("失败", failed),
        ("等待", 60 - completed - running - failed),
        ("数据集", ", ".join(DATASETS)),
        ("M", "1, 5, 10, 15, floor(sqrt(d))"),
        ("Seeds", "14, 42, 81"),
        ("Epochs / batch", "500 / 128"),
        ("固定设置", "delta=1e-3, scheduled momentum"),
        ("步长", json.dumps(STEP)),
        ("说明", "M=5 写入原配置行；公开 validation 固定作为测试集。"),
    ]
    for row in rows:
        progress.append(row)
    progress.column_dimensions["A"].width = 24
    progress.column_dimensions["B"].width = 90

    tmp = WORKBOOK.with_suffix(".tmp.xlsx")
    workbook.save(tmp)
    os.replace(tmp, WORKBOOK)
    print(f"{now_text()} status={counts} -> {WORKBOOK}")
    return counts


core.refresh = refresh


def session_exists(name: str) -> bool:
    return subprocess.run(
        ["tmux", "has-session", "-t", name],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    ).returncode == 0


def launch() -> None:
    build()
    refresh()
    for gpu in range(8):
        name = f"auc_highdim_gpu{gpu}"
        if not session_exists(name):
            command = f"cd {ROOT} && {sys.executable} {Path(__file__).resolve()} worker --gpu {gpu}"
            subprocess.run(["tmux", "new-session", "-d", "-s", name, command], check=True)
    report = "auc_highdim_report"
    if not session_exists(report):
        command = f"cd {ROOT} && {sys.executable} {Path(__file__).resolve()} reporter"
        subprocess.run(["tmux", "new-session", "-d", "-s", report, command], check=True)
    print("Launched auc_highdim_gpu0..7 and auc_highdim_report")


def reporter() -> None:
    while True:
        counts = refresh()
        runs = build()
        terminal = counts.get("completed", 0) + counts.get("failed", 0)
        if terminal >= len(runs):
            return
        for gpu in range(8):
            unfinished = any(
                run["gpu"] == gpu and load_state(run).get("status") != "completed"
                for run in runs
            )
            name = f"auc_highdim_gpu{gpu}"
            if unfinished and not session_exists(name):
                command = f"cd {ROOT} && {sys.executable} {Path(__file__).resolve()} worker --gpu {gpu}"
                subprocess.run(["tmux", "new-session", "-d", "-s", name, command], check=True)
        time.sleep(120)


def status() -> None:
    runs = build()
    counts = {}
    for run in runs:
        value = load_state(run).get("status", "pending")
        counts[value] = counts.get(value, 0) + 1
    print(json.dumps(counts, ensure_ascii=False))
    for gpu in range(8):
        name = f"auc_highdim_gpu{gpu}"
        print(name, "alive" if session_exists(name) else "stopped")


def main() -> None:
    parser = argparse.ArgumentParser()
    subparsers = parser.add_subparsers(dest="command", required=True)
    subparsers.add_parser("build").add_argument("--force", action="store_true")
    worker_parser = subparsers.add_parser("worker")
    worker_parser.add_argument("--gpu", type=int, choices=range(8), required=True)
    subparsers.add_parser("refresh")
    subparsers.add_parser("launch")
    subparsers.add_parser("reporter")
    subparsers.add_parser("status")
    args = parser.parse_args()

    if args.command == "build":
        build(args.force)
    elif args.command == "worker":
        core.worker(args.gpu)
    elif args.command == "refresh":
        refresh()
    elif args.command == "launch":
        launch()
    elif args.command == "reporter":
        reporter()
    elif args.command == "status":
        status()


if __name__ == "__main__":
    main()
